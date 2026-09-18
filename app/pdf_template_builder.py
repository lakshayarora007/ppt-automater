"""Build a PowerPoint from a PDF template with audit content and screenshots."""
from __future__ import annotations

import json
from pathlib import Path
from typing import Any

import pymupdf
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

WIDE, HIGH = Inches(13.333), Inches(7.5)
BLACK = RGBColor(10, 10, 10)
WHITE = RGBColor(255, 255, 255)
RED = RGBColor(220, 35, 35)
SOFT_RED = RGBColor(255, 157, 149)


def _shape(slide, x: float, y: float, w: float, h: float, color: RGBColor):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape


def _text(slide, value: str, x: float, y: float, w: float, h: float, size: float = 12, bold: bool = False, color: RGBColor = BLACK, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.vertical_anchor = MSO_ANCHOR.TOP
    paragraph = frame.paragraphs[0]
    paragraph.text = value
    paragraph.alignment = align
    paragraph.font.name = "Arial"
    paragraph.font.size = Pt(size)
    paragraph.font.bold = bold
    paragraph.font.color.rgb = color
    return box


def _render_template_pages(pdf_path: Path, assets_dir: Path) -> list[Path]:
    pdf = pymupdf.open(str(pdf_path))
    render_dir = assets_dir / "template-pages"
    render_dir.mkdir(parents=True, exist_ok=True)
    rendered = []
    for index, page in enumerate(pdf, start=1):
        destination = render_dir / f"{index:02d}.png"
        if not destination.exists():
            page.get_pixmap(matrix=pymupdf.Matrix(1.35, 1.35), alpha=False).save(str(destination))
        rendered.append(destination)
    return rendered


def _page_map(audit: dict[str, Any]) -> dict[str, list[dict[str, Any]]]:
    mapped: dict[str, list[dict[str, Any]]] = {}
    for page in audit.get("pages", []):
        mapped.setdefault(page.get("page_name", "Store Page"), []).append(page)
    return mapped


def _add_screenshot(slide, path: str | None, x: float, y: float, w: float = 2.5, h: float = 5.3):
    if path and Path(path).is_file():
        slide.shapes.add_picture(path, Inches(x), Inches(y), width=Inches(w), height=Inches(h))


def _replace_analysis(slide, title: str, issues: list[str], impact: str, screenshot: str | None):
    _shape(slide, 0.45, 1.35, 5.1, 4.9, WHITE)
    _text(slide, title, 0.55, 0.62, 5.0, 0.45, 25, True)
    _text(slide, "Current State", 0.8, 1.42, 2.2, 0.3, 13, True, RED)
    bullets = "\n".join(f"• {issue}" for issue in issues[:5])
    _text(slide, bullets or "• Validate the rendered experience before implementation.", 0.8, 1.85, 4.3, 2.3, 12)
    _shape(slide, 0.75, 5.1, 4.55, 0.78, SOFT_RED)
    _text(slide, impact, 0.95, 5.28, 4.15, 0.42, 11, False, BLACK, PP_ALIGN.CENTER)
    _shape(slide, 6.55, 0.82, 6.25, 5.85, RGBColor(242, 243, 243))
    _text(slide, "Our Brand", 7.2, 0.9, 2.3, 0.3, 12, True, RGBColor(45, 67, 95), PP_ALIGN.CENTER)
    _add_screenshot(slide, screenshot, 7.2, 1.45)


def _new_audit_slide(prs, brand: str, page: dict[str, Any], number: int):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _shape(slide, 0, 0, 13.333, 7.5, WHITE)
    _shape(slide, 0, 0, 13.333, 0.18, RGBColor(139, 205, 48))
    _text(slide, f"{brand.upper()}  /  MOBILE CRO AUDIT", 0.55, 0.42, 5.5, 0.25, 9, True, RGBColor(75, 75, 75))
    _text(slide, page.get("page_name", "Store Page"), 0.55, 0.82, 6.2, 0.5, 27, True)
    _text(slide, page.get("url", ""), 0.58, 1.38, 7.2, 0.25, 9, False, RGBColor(105, 105, 105))
    _shape(slide, 0.55, 1.88, 3.35, 4.85, RGBColor(247, 247, 247))
    _text(slide, "AUDIT POINTS", 0.8, 2.15, 2.2, 0.25, 10, True, RED)
    issues = page.get("issues", []) or [page.get("summary", "Validate the rendered experience before implementation.")]
    _text(slide, "\n".join(f"• {issue}" for issue in issues[:5]), 0.8, 2.62, 2.8, 2.35, 12)
    _shape(slide, 0.8, 5.35, 2.85, 0.85, SOFT_RED)
    _text(slide, "Fix the highest-friction\nstep before adding persuasion.", 0.95, 5.58, 2.55, 0.42, 10, True, BLACK, PP_ALIGN.CENTER)
    _text(slide, "CURRENT PLAYWRIGHT SCREEN", 4.35, 1.88, 2.8, 0.25, 9, True, RGBColor(80, 80, 80))
    _text(slide, "PROPOSED SCREEN", 8.75, 1.88, 2.2, 0.25, 9, True, RGBColor(80, 80, 80))
    _add_screenshot(slide, page.get("screenshot_path"), 4.45, 2.25, 3.35, 4.9)
    _add_screenshot(slide, page.get("proposed_screenshot_path"), 8.85, 2.25, 3.35, 4.9)
    _text(slide, f"{number:02d}", 12.3, 7.08, 0.45, 0.2, 9, True, RGBColor(100, 100, 100), PP_ALIGN.RIGHT)
    return slide


def _new_recommendation_slide(prs, brand: str, audit: dict[str, Any], title: str, items: list[str], number: int):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _shape(slide, 0, 0, 13.333, 7.5, WHITE)
    _shape(slide, 0, 0, 13.333, 0.18, RGBColor(139, 205, 48))
    _text(slide, f"{brand.upper()}  /  RECOMMENDED DIRECTION", 0.55, 0.42, 6.5, 0.25, 9, True, RGBColor(75, 75, 75))
    _text(slide, title, 0.55, 0.9, 10.5, 0.55, 29, True)
    _shape(slide, 0.65, 1.85, 12.0, 4.8, RGBColor(247, 247, 247))
    _text(slide, "PRIORITIES FROM THE AUDIT", 0.95, 2.15, 4.0, 0.25, 10, True, RED)
    _text(slide, "\n".join(f"0{index + 1}  {item}" for index, item in enumerate(items[:5])), 0.95, 2.7, 10.6, 2.7, 17, True)
    _text(slide, f"{number:02d}", 12.3, 7.08, 0.45, 0.2, 9, True, RGBColor(100, 100, 100), PP_ALIGN.RIGHT)


def build_pdf_template_deck(template_pdf: str, audit_path: str, output_path: str) -> None:
    template = Path(template_pdf)
    audit_file = Path(audit_path)
    output = Path(output_path)
    audit = json.loads(audit_file.read_text(encoding="utf-8"))
    pages = _page_map(audit)
    rendered = _render_template_pages(template, output.parent / f"{output.stem}-assets")

    prs = Presentation()
    prs.slide_width, prs.slide_height = WIDE, HIGH
    blank = prs.slide_layouts[6]
    # Keep the opening template pages, then insert the new audit section in the
    # same position as the removed brand-specific PDF pages.
    for image in rendered[:16]:
        slide = prs.slides.add_slide(blank)
        slide.shapes.add_picture(str(image), 0, 0, width=WIDE, height=HIGH)

    cover = prs.slides[0]
    _shape(cover, 0.55, 2.75, 8.6, 2.0, WHITE)
    _text(cover, "2026", 0.58, 2.82, 1.2, 0.3, 14)
    _text(cover, "CRO + AOV", 0.55, 3.35, 7.5, 0.7, 42, True)
    _text(cover, f"{audit.get('brand_name', 'Brand')} Growth Audit", 0.58, 4.08, 8.5, 0.5, 25, True)
    _text(cover, "Mobile conversion and basket-growth opportunities", 0.6, 4.68, 7.8, 0.35, 15)

    # Insert one clean audit slide for every crawled screenshot-backed page.
    page_number = 17
    for page in audit.get("pages", []):
        if page.get("screenshot_path"):
            _new_audit_slide(prs, audit.get("brand_name", "Brand"), page, page_number)
            page_number += 1

    recommendations = [item.get("title", "") for item in audit.get("recommendations", []) if item.get("title")]
    aov = [item.get("title", "") for item in audit.get("aov_opportunities", []) if item.get("title")]
    _new_recommendation_slide(prs, audit.get("brand_name", "Brand"), audit, "Make the mobile decision easier", recommendations, page_number)
    _new_recommendation_slide(prs, audit.get("brand_name", "Brand"), audit, "Grow basket value with relevance", aov, page_number + 1)

    # PDF pages 30-36 are also brand-specific audit/benchmark content. Keep
    # only the reusable company capability and testimonial pages after them.
    for image in rendered[36:]:
        slide = prs.slides.add_slide(blank)
        slide.shapes.add_picture(str(image), 0, 0, width=WIDE, height=HIGH)

    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(output))
