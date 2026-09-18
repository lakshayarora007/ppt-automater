from __future__ import annotations

from pptx import Presentation
from pptx.util import Inches, Pt
from pathlib import Path


def _add_title(slide, title: str, subtitle: str = ""):
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12.0), Inches(0.7))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(28)
    p.font.bold = True

    if subtitle:
        sub_box = slide.shapes.add_textbox(Inches(0.6), Inches(1.0), Inches(12.0), Inches(0.5))
        tf2 = sub_box.text_frame
        p2 = tf2.paragraphs[0]
        p2.text = subtitle
        p2.font.size = Pt(14)


def _add_bullets(slide, title: str, lines: list[str], left: float = 0.8, top: float = 1.4, width: float = 11.5, height: float = 5.2):
    title_box = slide.shapes.add_textbox(Inches(left), Inches(0.9), Inches(width), Inches(0.5))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(22)
    p.font.bold = True

    body_box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf2 = body_box.text_frame
    tf2.word_wrap = True
    for idx, line in enumerate(lines):
        p2 = tf2.paragraphs[0] if idx == 0 else tf2.add_paragraph()
        p2.text = line
        p2.level = 0
        p2.font.size = Pt(18)


def build_pptx_from_audit(data: dict, output_path: str):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Title slide
    title_slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_title(title_slide, data.get("brand_name", "CRO Audit"), f"Mobile CRO + AOV Audit\n{data.get('website_url', '')}")

    # Executive summary
    summary_slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_bullets(summary_slide, "Executive Summary", [data.get("executive_summary", "")])

    # Key findings
    findings_lines = []
    for page in data.get("pages", [])[:4]:
        findings_lines.append(f"{page.get('page_name', 'Page')} — {page.get('priority', 'P2')}: {page.get('summary', '')}")
    findings_slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_bullets(findings_slide, "Key Findings", findings_lines)

    # One evidence slide per crawled page. These are intentionally only created
    # when a real Playwright screenshot exists; heuristic fallback reports do
    # not fabricate visual evidence.
    for page in data.get("pages", [])[:8]:
        screenshot = page.get("screenshot_path")
        if not screenshot or not Path(screenshot).is_file():
            continue
        evidence_slide = prs.slides.add_slide(prs.slide_layouts[6])
        _add_title(evidence_slide, page.get("page_name", "Mobile Page"), page.get("priority", "P2"))
        issue_lines = page.get("issues", []) or [page.get("summary", "")]
        _add_bullets(evidence_slide, "Observed mobile evidence", issue_lines, left=0.55, top=1.65, width=5.4, height=4.8)
        proposed = page.get("proposed_screenshot_path")
        if proposed and Path(proposed).is_file():
            _add_title(evidence_slide, page.get("page_name", "Mobile Page"), "CURRENT → PROPOSED")
            evidence_slide.shapes.add_picture(screenshot, Inches(5.7), Inches(1.25), height=Inches(5.7))
            evidence_slide.shapes.add_picture(proposed, Inches(9.45), Inches(1.25), height=Inches(5.7))
        else:
            evidence_slide.shapes.add_picture(screenshot, Inches(7.1), Inches(1.1), height=Inches(5.95))

    # AOV opportunities
    aov_lines = []
    for opp in data.get("aov_opportunities", [])[:4]:
        aov_lines.append(f"{opp.get('title', '')} ({opp.get('placement', '')}) — {opp.get('description', '')}")
    aov_slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_bullets(aov_slide, "AOV Opportunities", aov_lines)

    # Recommendations
    recommendation_lines = []
    for rec in data.get("recommendations", [])[:5]:
        recommendation_lines.append(f"{rec.get('title', '')} — {rec.get('priority', 'P2')}: {rec.get('detail', '')}")
    rec_slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_bullets(rec_slide, "Priority Recommendations", recommendation_lines)

    # Deck outline
    outline_lines = [f"{idx + 1}. {item}" for idx, item in enumerate(data.get("deck_outline", []))]
    outline_slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_bullets(outline_slide, "Deck Outline", outline_lines)

    prs.save(output_path)
    print(f"Saved PPT: {output_path}")
